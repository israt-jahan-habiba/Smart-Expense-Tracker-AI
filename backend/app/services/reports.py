import io
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer


def monthly_xlsx(user, month, expenses, by_cat) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    total = sum(by_cat.values())
    inc = user.monthly_income_minor / 100
    for r in [("Report", month), ("Income", inc), ("Expenses", total), ("Savings", inc - total), (), ("Category", "Spent")] + list(by_cat.items()):
        ws.append(list(r))
    ws["A6"].font = ws["B6"].font = Font(bold=True)
    ex = wb.create_sheet("Expenses")
    ex.append(["Date", "Category", "Description", "Payment", "Amount"])
    for e in expenses:
        ex.append([e["date"], e["category"], e["description"], e["payment_method"], e["amount"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def monthly_pdf(user, month, expenses, by_cat) -> bytes:
    buf = io.BytesIO()
    st = getSampleStyleSheet()
    total, inc = sum(by_cat.values()), user.monthly_income_minor / 100
    cur = user.currency
    story = [Paragraph(f"Financial Report - {month}", st["Title"]), Paragraph(f"{user.name}", st["Normal"]), Spacer(1, 12),
             Paragraph(f"Income: {inc:,.2f} {cur} &nbsp; Expenses: {total:,.2f} {cur} &nbsp; Savings: {inc - total:,.2f} {cur}", st["Normal"]), Spacer(1, 12)]
    t = Table([["Category", "Spent"]] + [[k, f"{v:,.2f}"] for k, v in by_cat.items()])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14213D")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("ALIGN", (1, 0), (1, -1), "RIGHT")]))
    story.append(t)
    story += [Spacer(1, 16), Paragraph("Transactions", st["Heading2"])]
    t2 = Table([["Date", "Category", "Description", "Amount"]] + [[e["date"], e["category"], (e["description"] or "")[:40], f"{e['amount']:,.2f}"] for e in expenses[:200]], repeatRows=1)
    t2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECF2")), ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    story.append(t2)
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()
